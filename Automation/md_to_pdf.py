import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx2pdf import convert
import pypandoc

def ensure_pandoc_installed():
    """Ensures pandoc is available, downloading it automatically if missing."""
    try:
        pypandoc.get_pandoc_version()
    except (OSError, RuntimeError):
        print("Pandoc binary not found. Downloading Pandoc automatically...")
        try:
            pypandoc.download_pandoc()
            print("Pandoc downloaded and installed successfully!")
        except Exception as e:
            print(f"Failed to automatically download Pandoc: {e}")
            print("Please install it manually via: winget install JohnMacFarlane.Pandoc")
            exit(1)

def apply_custom_formatting(doc_path):
    """Applies Calibri font (10pt) and 0.5-inch narrow margins to a docx file."""
    doc = Document(doc_path)

    # 1. Set margins to 0.5 inches (narrow)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # 2. Apply font formatting to paragraphs and headings
    for paragraph in doc.paragraphs:
        is_heading = paragraph.style.name.startswith("Heading")
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if is_heading:
                run.font.bold = True

    # 3. Apply font formatting to tables (if any exist)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(10)

    doc.save(doc_path)

def process_markdown_directory(root_dir_path, output_pdf_name="combined_output.pdf"):
    # Ensure pandoc is ready before starting
    ensure_pandoc_installed()

    root_path = Path(root_dir_path)
    if not root_path.exists():
        print(f"Error: The directory '{root_dir_path}' does not exist.")
        return

    # Find all .md files recursively
    md_files = sorted(list(root_path.rglob("*.md")))
    
    if not md_files:
        print(f"No Markdown (.md) files found in '{root_dir_path}'.")
        return

    print(f"Found {len(md_files)} Markdown files. Starting conversion...")

    generated_docx_files = []

    # Step 1, 2, 3, 5, 6: Convert each .md to .docx and apply formatting
    for md_file in md_files:
        docx_file = md_file.with_suffix(".docx")
        print(f"Converting: {md_file.relative_to(root_path)} -> .docx")
        
        # Convert Markdown to Docx using pypandoc
        pypandoc.convert_file(str(md_file), 'docx', outputfile=str(docx_file))
        
        # Apply margins, font (Calibri), and size (10)
        apply_custom_formatting(docx_file)
        generated_docx_files.append(docx_file)

    # Step 4: Combine all generated docx files into a single master docx
    print("Combining all .docx files into a single document...")
    master_docx_path = root_path / "master_combined.docx"
    
    if generated_docx_files:
        master_doc = Document(generated_docx_files[0])
        
        for docx_file in generated_docx_files[1:]:
            master_doc.add_page_break()  # Add a page break between different files
            sub_doc = Document(docx_file)
            for element in sub_doc.element.body:
                master_doc.element.body.append(element)
                
        master_doc.save(master_docx_path)

        # Convert the master docx to a single PDF
        output_pdf_path = root_path / output_pdf_name
        print(f"Converting combined document to PDF: {output_pdf_path}")
        convert(str(master_docx_path), str(output_pdf_path))
        
        print(f"Success! Final print-ready PDF generated at:\n{output_pdf_path}")

if __name__ == "__main__":
    target_directory = input("Enter the path to the directory containing .md files: ").strip()
    process_markdown_directory(target_directory)